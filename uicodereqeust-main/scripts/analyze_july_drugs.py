import openpyxl
import re
from collections import Counter

wb = openpyxl.load_workbook(r'C:\Users\WINDOWS\Downloads\july ao.xlsx')
sheet = wb.active
treatments = []
for r in range(1, sheet.max_row + 1):
    val = sheet.cell(r, 7).value
    if val:
        treatments.append(str(val))

with open('supabase/migrations/20260620090000_fresh_nhia_2025_tariffs.sql', 'r', encoding='utf-8') as f:
    tariffs_text = f.read()

# All drugs in tariff
tariffs = re.findall(r"\('([^']+)',\s*'([^']+)',\s*'([^']+)',\s*([0-9.]+)\)", tariffs_text)
print(f'Total tariffs: {len(tariffs)}')

for drug in ['labetalol', 'metoprolol', 'donepezil', 'atorvastatin']:
    print(f'=== {drug} ===')
    matches = [t for t in tariffs if drug in t[1].lower()]
    for m in matches:
        print(m)

# Count clinical terms in treatments
all_text = ' '.join(treatments).lower()
cleaned = re.sub(r'\d+(?:\.\d+)?\s*(?:mg|g|mcg|ml|tabs?|caps?|each|nocte|daily|dly)', ' ', all_text)
words = re.findall(r'[a-z]{4,}', cleaned)
stopwords = {'daily', 'tabs', 'each', 'nocte', 'tablets', 'tablet', 'oral', 'dose', 'drug', 'referral', 'requested'}
counts = Counter(w for w in words if w not in stopwords)
print('\nTop 35 clinical terms in treatments:')
for w, c in counts.most_common(35):
    print(f'  {w}: {c}')
